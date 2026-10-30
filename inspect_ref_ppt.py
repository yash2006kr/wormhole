import pptx
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
import os

# Read the uploaded reference PPT
ref_path = r'C:\Users\Yashwanth KR\.gemini\antigravity\brain\58c06b0e-74fa-452a-a6a9-130c71f087d4\.user_uploaded\media_1792857960745_b53dd461.pptx'
prs = Presentation(ref_path)

print(f"=== REFERENCE PPT ===")
print(f"Slide size: {prs.slide_width.inches:.2f}\" x {prs.slide_height.inches:.2f}\"")
print(f"Total slides: {len(prs.slides)}")

for idx, slide in enumerate(prs.slides):
    print(f"\n=== SLIDE {idx+1} ===")
    for s_idx, shape in enumerate(slide.shapes):
        shape_type = shape.shape_type
        left = round(shape.left / 914400, 2)
        top = round(shape.top / 914400, 2)
        width = round(shape.width / 914400, 2)
        height = round(shape.height / 914400, 2)
        
        if shape.has_text_frame:
            text = shape.text_frame.text.replace('\n', ' | ')[:80]
            print(f"  [{s_idx}] TEXTBOX '{shape.name}': pos=({left},{top}) size=({width}x{height}) | '{text}'")
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                align = p.alignment
                for r in p.runs:
                    fname = r.font.name
                    fsize = r.font.size.pt if r.font.size else None
                    bold = r.font.bold
                    try:
                        color = str(r.font.color.rgb) if r.font.color and r.font.color.type else None
                    except:
                        color = None
                    print(f"       Run: '{r.text[:50]}' font={fname} size={fsize} bold={bold} color={color} align={align}")
        elif shape_type == 13:  # PICTURE
            print(f"  [{s_idx}] PICTURE '{shape.name}': pos=({left},{top}) size=({width}x{height})")
            try:
                img = shape.image
                print(f"       Image content_type={img.content_type}, blob_size={len(img.blob)} bytes")
                # Save image
                ext = img.content_type.split('/')[-1].replace('jpeg','jpg')
                img_path = f'ref_img_slide{idx+1}_{s_idx}.{ext}'
                with open(img_path, 'wb') as f:
                    f.write(img.blob)
                print(f"       Saved: {img_path}")
            except Exception as e:
                print(f"       Image error: {e}")
        else:
            print(f"  [{s_idx}] SHAPE '{shape.name}' type={shape_type}: pos=({left},{top}) size=({width}x{height})")
            # Check fill
            try:
                fill = shape.fill
                if fill.type is not None:
                    print(f"       Fill type: {fill.type}")
                    if hasattr(fill, 'fore_color') and fill.fore_color:
                        try:
                            print(f"       Fill color: {fill.fore_color.rgb}")
                        except:
                            pass
            except:
                pass
