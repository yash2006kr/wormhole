import pptx

prs = pptx.Presentation('AI_Avengers_4.pptx')
print(f"Total slides: {len(prs.slides)}")

for i, slide in enumerate(prs.slides):
    print(f"\n--- Slide {i+1} ---")
    for sh in slide.shapes:
        txt = ""
        if sh.has_text_frame:
            txt = sh.text_frame.text.replace("\n", " | ")
        print(f"Shape: {sh.name} (type={sh.shape_type}) pos=({sh.left/914400:.2f}, {sh.top/914400:.2f}) size=({sh.width/914400:.2f} x {sh.height/914400:.2f}) | text: {txt[:80]}")
