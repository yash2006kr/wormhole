import pptx
from pptx import Presentation

prs = Presentation('AfterHours_1.0_DeepGuard_Deepfake_Tracker.pptx')
print(f"Total slides: {len(prs.slides)}")

with open('deepguard_deck_check.txt', 'w', encoding='utf-8') as f:
    for idx, slide in enumerate(prs.slides):
        f.write(f"\n================ SLIDE {idx+1} ================\n")
        pics = [s for s in slide.shapes if s.shape_type == 13]
        f.write(f"Pictures ({len(pics)}):\n")
        for p in pics:
            f.write(f"  Pic: {p.name}, left={p.left/914400:.2f}in, top={p.top/914400:.2f}in, w={p.width/914400:.2f}in, h={p.height/914400:.2f}in\n")
        
        textboxes = [s for s in slide.shapes if s.has_text_frame]
        f.write(f"Text shapes ({len(textboxes)}):\n")
        for tb in textboxes:
            for p in tb.text_frame.paragraphs:
                txt = p.text.strip().replace('\n', ' ')
                if txt:
                    run_sizes = [r.font.size.pt for r in p.runs if r.font and r.font.size]
                    f.write(f"  [Txt]: run_sizes={run_sizes} | '{txt[:70]}'\n")

print("Saved deepguard_deck_check.txt!")
