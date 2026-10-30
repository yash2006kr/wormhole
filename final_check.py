from pptx import Presentation
prs = Presentation('AfterHours_1.0_DeepfakeDetection_DeepGuard.pptx')
with open('final_check.txt', 'w', encoding='utf-8') as f:
    f.write(f'Slides: {len(prs.slides)}\n')
    for i, slide in enumerate(prs.slides):
        pics = [s for s in slide.shapes if s.shape_type == 13]
        textboxes = [s for s in slide.shapes if s.has_text_frame]
        f.write(f'Slide {i+1}: {len(pics)} images, {len(textboxes)} text boxes\n')
        for tb in textboxes:
            for p in tb.text_frame.paragraphs:
                txt = p.text.strip()[:60]
                run_sizes = [r.font.size.pt for r in p.runs if r.font.size]
                if txt:
                    f.write(f'  P: {run_sizes} | {txt}\n')
print('Done — check final_check.txt')
